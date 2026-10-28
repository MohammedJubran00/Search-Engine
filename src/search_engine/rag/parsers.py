"""Secure text extraction for supported RAG document formats.

Why it exists: the RAG contracts define parsing but no implementation can
validate and extract PDF, DOCX, or UTF-8 text content. This module performs
CPU/blocking parser work only and is called through ``asyncio.to_thread`` by
the document service.
"""

import math
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document as load_docx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.search_engine.rag.contracts import ParsedDocument, ParsedPage

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
TXT_CONTENT_TYPE = "text/plain"


class DocumentParseError(ValueError):
    """Base class for safe client-facing parsing failures."""


class CorruptDocumentError(DocumentParseError):
    """Raised when bytes do not form the declared document type."""


class EmptyDocumentError(DocumentParseError):
    """Raised when a document contains no extractable text."""


class ScannedPdfError(DocumentParseError):
    """Raised when a PDF needs OCR, which is intentionally out of scope."""


class UnsupportedContentTypeError(DocumentParseError):
    """Raised when no Phase 2 parser supports the declared MIME type."""


class RagDocumentParser:
    """Validate and extract supported documents without performing OCR."""

    def __init__(
        self,
        *,
        max_docx_uncompressed_bytes: int,
        pdf_boilerplate_min_pages: int = 3,
        pdf_boilerplate_page_ratio: float = 0.6,
        pdf_boilerplate_margin_ratio: float = 0.08,
        pdf_boilerplate_max_line_characters: int = 200,
    ) -> None:
        _validate_pdf_boilerplate_policy(
            min_pages=pdf_boilerplate_min_pages,
            page_ratio=pdf_boilerplate_page_ratio,
            margin_ratio=pdf_boilerplate_margin_ratio,
            max_line_characters=pdf_boilerplate_max_line_characters,
        )
        self._max_docx_uncompressed_bytes = max_docx_uncompressed_bytes
        self._pdf_boilerplate_min_pages = pdf_boilerplate_min_pages
        self._pdf_boilerplate_page_ratio = pdf_boilerplate_page_ratio
        self._pdf_boilerplate_margin_ratio = pdf_boilerplate_margin_ratio
        self._pdf_boilerplate_max_line_characters = (
            pdf_boilerplate_max_line_characters
        )

    def parse(self, path: Path, content_type: str) -> ParsedDocument:
        if content_type == PDF_CONTENT_TYPE:
            return self._parse_pdf(path)
        if content_type == DOCX_CONTENT_TYPE:
            return self._parse_docx(path)
        if content_type == TXT_CONTENT_TYPE:
            return self._parse_txt(path)
        raise UnsupportedContentTypeError("Unsupported document content type.")

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        data = path.read_bytes()
        if not data.startswith(b"%PDF-"):
            raise CorruptDocumentError("File content is not a valid PDF.")

        try:
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.needs_pass:
                    raise CorruptDocumentError("Password-protected PDFs are unsupported.")

                extracted_pages = tuple(
                    _extract_pdf_page(
                        page,
                        page_number=index + 1,
                        margin_ratio=self._pdf_boilerplate_margin_ratio,
                        max_line_characters=self._pdf_boilerplate_max_line_characters,
                    )
                    for index, page in enumerate(document)
                )
                pages, removed_line_count, pattern_count = _remove_pdf_boilerplate(
                    extracted_pages,
                    min_pages=self._pdf_boilerplate_min_pages,
                    page_ratio=self._pdf_boilerplate_page_ratio,
                )
                metadata = _clean_metadata(document.metadata)
                metadata["page_count"] = document.page_count
                metadata["removed_boilerplate_lines"] = removed_line_count
                metadata["removed_boilerplate_patterns"] = pattern_count
        except DocumentParseError:
            raise
        except (pymupdf.FileDataError, RuntimeError, ValueError) as exc:
            raise CorruptDocumentError("PDF is corrupted or unreadable.") from exc

        text = "\n\n".join(page.text for page in pages)
        if not text.strip():
            raise ScannedPdfError(
                "PDF contains no extractable text; scanned or image-only PDFs require OCR."
            )
        return ParsedDocument(text=text, pages=pages, metadata=metadata)

    def _parse_docx(self, path: Path) -> ParsedDocument:
        _validate_docx_container(path, self._max_docx_uncompressed_bytes)
        try:
            document = load_docx(path)
            blocks: list[str] = []
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    blocks.append(item.text)
                elif isinstance(item, Table):
                    blocks.extend(
                        "\t".join(cell.text for cell in row.cells)
                        for row in item.rows
                    )
            text = "\n".join(blocks)
            metadata = _docx_metadata(document)
        except (BadZipFile, KeyError, PackageNotFoundError, ValueError) as exc:
            raise CorruptDocumentError("DOCX is corrupted or unreadable.") from exc

        if not text.strip():
            raise EmptyDocumentError("Document contains no extractable text.")
        return ParsedDocument(
            text=text,
            pages=(ParsedPage(text=text),),
            metadata=metadata,
        )

    def _parse_txt(self, path: Path) -> ParsedDocument:
        data = path.read_bytes()
        if b"\x00" in data:
            raise CorruptDocumentError("TXT file contains binary data.")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise CorruptDocumentError("TXT files must use UTF-8 encoding.") from exc

        if any(ord(character) < 32 and character not in "\t\n\r" for character in text):
            raise CorruptDocumentError("TXT file contains unsupported control bytes.")
        if not text.strip():
            raise EmptyDocumentError("Document contains no extractable text.")
        return ParsedDocument(
            text=text,
            pages=(ParsedPage(text=text),),
            metadata={"encoding": "utf-8"},
        )


@dataclass(frozen=True, slots=True)
class _ExtractedPdfPage:
    text: str
    page_number: int
    edge_fingerprints: frozenset[str]


_TERMINAL_PAGE_NUMBER = re.compile(r"(?<!\d)(\d+)\s*$")
_PAGE_NUMBER_DELIMITER = re.compile(r"(?:[•·|]|[-–—]|\bpage|\bp\.)\s*$")


def _extract_pdf_page(
    page: pymupdf.Page,
    *,
    page_number: int,
    margin_ratio: float,
    max_line_characters: int,
) -> _ExtractedPdfPage:
    edge_fingerprints: set[str] = set()
    page_height = float(page.rect.height)
    for block in page.get_text("blocks", sort=True):
        normalized_lines = [
            _normalize_pdf_line(line)
            for line in str(block[4]).splitlines()
            if line.strip()
        ]
        if len(normalized_lines) != 1:
            continue
        line = normalized_lines[0]
        if not line or len(line) > max_line_characters:
            continue
        top_edge = float(block[3]) <= page_height * margin_ratio
        bottom_edge = float(block[1]) >= page_height * (1.0 - margin_ratio)
        if top_edge or bottom_edge:
            edge_fingerprints.add(_pdf_line_fingerprint(line, page_number))
    return _ExtractedPdfPage(
        text=page.get_text("text", sort=True),
        page_number=page_number,
        edge_fingerprints=frozenset(edge_fingerprints),
    )


def _remove_pdf_boilerplate(
    pages: tuple[_ExtractedPdfPage, ...],
    *,
    min_pages: int,
    page_ratio: float,
) -> tuple[tuple[ParsedPage, ...], int, int]:
    pattern_pages: dict[str, set[int]] = defaultdict(set)
    for page in pages:
        for fingerprint in page.edge_fingerprints:
            pattern_pages[fingerprint].add(page.page_number)
    required_pages = max(min_pages, math.ceil(len(pages) * page_ratio))
    verified = {
        fingerprint
        for fingerprint, page_numbers in pattern_pages.items()
        if len(page_numbers) >= required_pages
    }

    cleaned_pages: list[ParsedPage] = []
    removed_line_count = 0
    for page in pages:
        kept_lines: list[str] = []
        for line in page.text.splitlines():
            normalized = _normalize_pdf_line(line)
            fingerprint = _pdf_line_fingerprint(normalized, page.page_number)
            if (
                normalized
                and fingerprint in verified
                and fingerprint in page.edge_fingerprints
            ):
                removed_line_count += 1
                continue
            kept_lines.append(line)
        text = "\n".join(kept_lines)
        if page.text.endswith("\n") and text:
            text += "\n"
        cleaned_pages.append(ParsedPage(text=text, page_number=page.page_number))
    return tuple(cleaned_pages), removed_line_count, len(verified)


def _normalize_pdf_line(line: str) -> str:
    return " ".join(line.split())


def _pdf_line_fingerprint(line: str, page_number: int) -> str:
    normalized = _normalize_pdf_line(line).casefold()
    match = _TERMINAL_PAGE_NUMBER.search(normalized)
    if match is not None and int(match.group(1)) == page_number:
        prefix = normalized[: match.start()].rstrip()
        if not prefix or _PAGE_NUMBER_DELIMITER.search(prefix):
            normalized = f"{prefix} <page>".strip()
    return normalized


def _validate_pdf_boilerplate_policy(
    *,
    min_pages: int,
    page_ratio: float,
    margin_ratio: float,
    max_line_characters: int,
) -> None:
    if min_pages < 3:
        raise ValueError("PDF boilerplate detection requires at least three pages.")
    if not 0.5 < page_ratio <= 1.0:
        raise ValueError("PDF boilerplate page ratio must be above 0.5 and at most 1.")
    if not 0.0 < margin_ratio <= 0.2:
        raise ValueError("PDF boilerplate margin ratio must be above 0 and at most 0.2.")
    if max_line_characters <= 0:
        raise ValueError("PDF boilerplate maximum line length must be positive.")


def _validate_docx_container(path: Path, max_uncompressed_bytes: int) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if sum(member.file_size for member in members) > max_uncompressed_bytes:
                raise CorruptDocumentError("DOCX expands beyond the safe size limit.")
            if any(member.flag_bits & 0x1 for member in members):
                raise CorruptDocumentError("Encrypted DOCX files are unsupported.")
            names = {member.filename for member in members}
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(names):
                raise CorruptDocumentError("File content is not a valid DOCX.")
            bad_member = archive.testzip()
            if bad_member is not None:
                raise CorruptDocumentError("DOCX is corrupted or unreadable.")
    except BadZipFile as exc:
        raise CorruptDocumentError("File content is not a valid DOCX.") from exc


def _clean_metadata(raw: dict[str, object] | None) -> dict[str, str | int]:
    if not raw:
        return {}
    return {
        key: str(value)
        for key, value in raw.items()
        if value not in (None, "")
    }


def _docx_metadata(document: DocxDocument) -> dict[str, str]:
    properties = document.core_properties
    values = {
        "title": properties.title,
        "author": properties.author,
        "subject": properties.subject,
        "language": properties.language,
    }
    return {key: value for key, value in values.items() if value}
