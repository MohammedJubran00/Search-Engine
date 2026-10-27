"""Phase 2 document parsing, storage, ownership, and cleanup tests."""

import asyncio
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from fastapi import UploadFile
from fastapi.testclient import TestClient
from sqlalchemy.exc import SQLAlchemyError
from starlette.datastructures import Headers

from src.search_engine.core.config import settings
from src.search_engine.models.document import Document
from src.search_engine.rag.parsers import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    TXT_CONTENT_TYPE,
    CorruptDocumentError,
    RagDocumentParser,
    ScannedPdfError,
)
from src.search_engine.rag.storage import (
    DocumentStorage,
    EmptyUploadError,
    FileTooLargeError,
)
from src.search_engine.schemas.document import DocumentResponse
from src.search_engine.services.document_service import (
    DocumentNotFoundError,
    DocumentPersistenceError,
    DocumentService,
    InvalidUploadError,
)
from tests.auth_helpers import auth_headers


class FakeSession:
    def __init__(self, *, fail_commit: bool = False) -> None:
        self.fail_commit = fail_commit
        self.commits = 0
        self.rollbacks = 0

    async def commit(self) -> None:
        self.commits += 1
        if self.fail_commit:
            raise SQLAlchemyError("forced commit failure")

    async def rollback(self) -> None:
        self.rollbacks += 1


class FakeDocumentRepository:
    def __init__(self) -> None:
        self.documents: dict[uuid.UUID, Document] = {}

    async def create(self, **fields: object) -> Document:
        now = datetime.now(timezone.utc)
        document = Document(
            id=uuid.uuid4(),
            processing_status="pending",
            created_at=now,
            updated_at=now,
            **fields,
        )
        self.documents[document.id] = document
        return document

    async def list_for_user(self, user_id: uuid.UUID) -> list[Document]:
        return [
            document
            for document in self.documents.values()
            if document.user_id == user_id
        ]

    async def get_by_id_for_user(
        self,
        document_id: uuid.UUID,
        user_id: uuid.UUID,
    ) -> Document | None:
        document = self.documents.get(document_id)
        if document is None or document.user_id != user_id:
            return None
        return document

    async def set_status(
        self,
        document: Document,
        status: str,
        *,
        error_message: str | None = None,
    ) -> None:
        document.processing_status = status
        document.error_message = error_message
        document.updated_at = datetime.now(timezone.utc)

    async def delete(self, document: Document) -> None:
        del self.documents[document.id]


def _upload(filename: str, content_type: str, data: bytes) -> UploadFile:
    return UploadFile(
        file=BytesIO(data),
        filename=filename,
        size=len(data),
        headers=Headers({"content-type": content_type}),
    )


def _pdf_bytes(*page_texts: str) -> bytes:
    document = pymupdf.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def _docx_bytes(text: str) -> bytes:
    buffer = BytesIO()
    document = DocxDocument()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _service(tmp_path: Path, *, fail_commit: bool = False):
    session = FakeSession(fail_commit=fail_commit)
    repository = FakeDocumentRepository()
    service = DocumentService(
        session,  # type: ignore[arg-type]
        repository=repository,  # type: ignore[arg-type]
        storage=DocumentStorage(tmp_path),
    )
    return service, session, repository


@pytest.mark.parametrize(
    ("filename", "content_type", "data", "expected_text"),
    [
        ("report.pdf", PDF_CONTENT_TYPE, _pdf_bytes("Page one"), "Page one"),
        (
            "report.docx",
            DOCX_CONTENT_TYPE,
            _docx_bytes("مرحبا بالعالم — résumé"),
            "مرحبا بالعالم",
        ),
        (
            "report.txt",
            TXT_CONTENT_TYPE,
            "نص عربي وUnicode ✓".encode(),
            "نص عربي",
        ),
    ],
    ids=["pdf", "docx", "txt-unicode"],
)
def test_supported_uploads_are_parsed_and_stored(
    tmp_path: Path,
    filename: str,
    content_type: str,
    data: bytes,
    expected_text: str,
) -> None:
    async def scenario() -> None:
        service, session, _ = _service(tmp_path)
        user_id = uuid.uuid4()
        result = await service.upload(
            user_id=user_id,
            upload=_upload(filename, content_type, data),
        )

        assert result.document.user_id == user_id
        assert result.document.processing_status == "completed"
        assert result.document.storage_path != filename
        assert expected_text in result.parsed.text
        assert (tmp_path / result.document.storage_path).read_bytes() == data
        assert session.commits == 1

    asyncio.run(scenario())


def test_pdf_parser_preserves_page_numbers_and_blank_pages(tmp_path: Path) -> None:
    path = tmp_path / "multi.pdf"
    path.write_bytes(_pdf_bytes("First page", "", "Third page"))

    parsed = RagDocumentParser(
        max_docx_uncompressed_bytes=settings.rag_max_docx_uncompressed_bytes
    ).parse(path, PDF_CONTENT_TYPE)

    assert [page.page_number for page in parsed.pages] == [1, 2, 3]
    assert "First page" in parsed.pages[0].text
    assert parsed.pages[1].text == ""
    assert "Third page" in parsed.pages[2].text
    assert parsed.metadata["page_count"] == 3


def test_docx_expanded_size_limit_is_enforced(tmp_path: Path) -> None:
    path = tmp_path / "large.docx"
    path.write_bytes(_docx_bytes("valid but larger than ten expanded bytes"))

    with pytest.raises(CorruptDocumentError, match="safe size limit"):
        RagDocumentParser(max_docx_uncompressed_bytes=10).parse(
            path,
            DOCX_CONTENT_TYPE,
        )


@pytest.mark.parametrize(
    ("filename", "content_type"),
    [
        ("malware.exe", "application/octet-stream"),
        ("fake.pdf", TXT_CONTENT_TYPE),
        ("../escape.txt", TXT_CONTENT_TYPE),
        (r"C:\escape.txt", TXT_CONTENT_TYPE),
    ],
)
def test_unsafe_or_mismatched_upload_metadata_is_rejected(
    tmp_path: Path,
    filename: str,
    content_type: str,
) -> None:
    async def scenario() -> None:
        service, _, _ = _service(tmp_path)
        with pytest.raises(InvalidUploadError):
            await service.upload(
                user_id=uuid.uuid4(),
                upload=_upload(filename, content_type, b"safe text"),
            )

    asyncio.run(scenario())


def test_empty_and_oversized_uploads_are_rejected_and_cleaned(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def scenario() -> None:
        service, _, _ = _service(tmp_path)
        with pytest.raises(EmptyUploadError):
            await service.upload(
                user_id=uuid.uuid4(),
                upload=_upload("empty.txt", TXT_CONTENT_TYPE, b""),
            )

        monkeypatch.setattr(settings, "rag_max_upload_bytes", 4)
        with pytest.raises(FileTooLargeError):
            await service.upload(
                user_id=uuid.uuid4(),
                upload=_upload("large.txt", TXT_CONTENT_TYPE, b"12345"),
            )
        assert list(tmp_path.iterdir()) == []

    asyncio.run(scenario())


@pytest.mark.parametrize(
    ("filename", "content_type", "data", "error_type"),
    [
        ("broken.pdf", PDF_CONTENT_TYPE, b"%PDF-broken", CorruptDocumentError),
        ("broken.docx", DOCX_CONTENT_TYPE, b"PK\x03\x04broken", CorruptDocumentError),
        ("binary.txt", TXT_CONTENT_TYPE, b"text\x00binary", CorruptDocumentError),
        ("scanned.pdf", PDF_CONTENT_TYPE, _pdf_bytes(""), ScannedPdfError),
    ],
    ids=["corrupt-pdf", "corrupt-docx", "binary-txt", "image-only-pdf"],
)
def test_corrupt_binary_and_image_only_documents_are_cleaned(
    tmp_path: Path,
    filename: str,
    content_type: str,
    data: bytes,
    error_type: type[Exception],
) -> None:
    async def scenario() -> None:
        service, session, _ = _service(tmp_path)
        with pytest.raises(error_type):
            await service.upload(
                user_id=uuid.uuid4(),
                upload=_upload(filename, content_type, data),
            )
        assert session.rollbacks == 1
        assert list(tmp_path.iterdir()) == []

    asyncio.run(scenario())


def test_ownership_listing_retrieval_and_deletion_cleanup(tmp_path: Path) -> None:
    async def scenario() -> None:
        service, _, _ = _service(tmp_path)
        owner_id = uuid.uuid4()
        other_id = uuid.uuid4()
        result = await service.upload(
            user_id=owner_id,
            upload=_upload("owned.txt", TXT_CONTENT_TYPE, b"owned text"),
        )
        document_id = result.document.id

        assert [item.id for item in await service.list_for_user(owner_id)] == [document_id]
        assert await service.list_for_user(other_id) == []
        assert (await service.get_for_user(
            document_id=document_id,
            user_id=owner_id,
        )).id == document_id
        with pytest.raises(DocumentNotFoundError):
            await service.get_for_user(document_id=document_id, user_id=other_id)

        await service.delete_for_user(document_id=document_id, user_id=owner_id)
        assert list(tmp_path.iterdir()) == []
        with pytest.raises(DocumentNotFoundError):
            await service.delete_for_user(document_id=document_id, user_id=owner_id)

    asyncio.run(scenario())


def test_database_failure_rolls_back_and_removes_stored_file(tmp_path: Path) -> None:
    async def scenario() -> None:
        service, session, _ = _service(tmp_path, fail_commit=True)
        with pytest.raises(DocumentPersistenceError):
            await service.upload(
                user_id=uuid.uuid4(),
                upload=_upload("db-fail.txt", TXT_CONTENT_TYPE, b"valid text"),
            )
        assert session.rollbacks == 1
        assert list(tmp_path.iterdir()) == []

    asyncio.run(scenario())


def test_public_document_schema_does_not_expose_storage_path() -> None:
    now = datetime.now(timezone.utc)
    document = Document(
        id=uuid.uuid4(),
        user_id=uuid.uuid4(),
        original_filename="private.txt",
        storage_path="secret-reference.txt",
        content_type=TXT_CONTENT_TYPE,
        file_size=12,
        processing_status="completed",
        created_at=now,
        updated_at=now,
    )

    payload = DocumentResponse.model_validate(document).model_dump()

    assert "storage_path" not in payload
    assert "error_message" not in payload


def test_upload_requires_authentication(client: TestClient) -> None:
    response = client.post(
        "/api/documents/upload",
        files={"file": ("private.txt", b"private", TXT_CONTENT_TYPE)},
    )

    assert response.status_code == 401


def test_document_endpoints_enforce_ownership_and_cleanup(
    client: TestClient,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "rag_storage_dir", tmp_path)
    owner_headers = auth_headers(client)
    upload = client.post(
        "/api/documents/upload",
        headers=owner_headers,
        files={
            "file": (
                "مرحبا.txt",
                "محتوى عربي آمن".encode(),
                TXT_CONTENT_TYPE,
            )
        },
    )

    assert upload.status_code == 201
    body = upload.json()
    document_id = body["id"]
    assert body["processing_status"] == "completed"
    assert body["original_filename"] == "مرحبا.txt"
    assert "storage_path" not in body
    assert len(list(tmp_path.iterdir())) == 1

    listing = client.get("/api/documents", headers=owner_headers)
    assert listing.status_code == 200
    assert document_id in {item["id"] for item in listing.json()}

    retrieved = client.get(
        f"/api/documents/{document_id}",
        headers=owner_headers,
    )
    assert retrieved.status_code == 200
    assert retrieved.json()["id"] == document_id

    other_headers = auth_headers(client)
    assert client.get(
        f"/api/documents/{document_id}",
        headers=other_headers,
    ).status_code == 404
    assert client.delete(
        f"/api/documents/{document_id}",
        headers=other_headers,
    ).status_code == 404

    deleted = client.delete(
        f"/api/documents/{document_id}",
        headers=owner_headers,
    )
    assert deleted.status_code == 204
    assert list(tmp_path.iterdir()) == []
    assert client.delete(
        f"/api/documents/{document_id}",
        headers=owner_headers,
    ).status_code == 404
